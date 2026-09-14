"""plan_8_2 section 1.2: one decoder for every table bioAF reads, whoever reads it.

Acquisition, supplement inspection and the consistency worker each decoded bytes their own way. The
SAMD1 undifferentiated author table is gzip-compressed UTF-16 with a byte-order mark: acquisition called
it binary, and the consistency worker decoded it as UTF-8 with replacement, producing NUL characters
that PostgreSQL rejected in JSONB. That failure rolled back the check queue's transaction every tick.

**Order.** Compression is recognized first (gzip, bzip2, xz, a single-member zip), then the enclosed
format from its magic bytes (a workbook, legacy or current), then a byte-order mark (UTF-32 and UTF-16 in
either order, UTF-8), and only then strict UTF-8. The extension is a hint and the bytes are the evidence.

**Strict.** Nothing is decoded with replacement. An encoding bioAF cannot establish (invalid UTF-8 with
no mark, UTF-16 with no mark) is a typed unresolved outcome, never a guess. A NUL character in decoded
text invalidates the table's interpretation: it is never stripped from scientific content, because a
table that needs repair is not the table the authors published.

**Distinct facts.** A file that arrived and could not be interpreted says so. It is never reported as a
file that was not obtained. The decoded artifact records the source checksum, compression, encoding,
format and decoder version, so every consumer can tell that it read the same content.
"""

from __future__ import annotations

import bz2
import hashlib
import io
import lzma
import zipfile
import zlib
from dataclasses import dataclass

# Bumped whenever the decoded text for some bytes can change, so a decoded artifact is reused only by the
# decoder that produced it.
DECODER_VERSION = 1

DECODED = "decoded"
UNRESOLVED = "unresolved"

# Why a payload was not decoded. Each is a typed outcome a consumer can act on without parsing words.
EMPTY = "empty"
BINARY = "binary"
ENCODING = "encoding"
NUL = "nul"
CORRUPT = "corrupt_compression"
TOO_LARGE = "too_large"
UNREADABLE_WORKBOOK = "unreadable_workbook"
ARCHIVE = "archive"
REASON_KINDS = (EMPTY, BINARY, ENCODING, NUL, CORRUPT, TOO_LARGE, UNREADABLE_WORKBOOK, ARCHIVE)

# The most a payload may decompress to when the caller states no smaller limit.
DEFAULT_MAX_DECOMPRESSED_BYTES = 2 * 1024**3

_GZIP = b"\x1f\x8b"
_BZIP2 = b"BZh"
_XZ = b"\xfd7zXZ\x00"
_ZIP = b"PK\x03\x04"
_OLE2 = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

# Longest first: a UTF-32 little-endian mark begins with the UTF-16 little-endian one.
_MARKS = (
    (b"\xff\xfe\x00\x00", "utf-32-le"),
    (b"\x00\x00\xfe\xff", "utf-32-be"),
    (b"\xef\xbb\xbf", "utf-8-sig"),
    (b"\xff\xfe", "utf-16-le"),
    (b"\xfe\xff", "utf-16-be"),
)

# plan_8_2 label, pending the owner's sign-off.
ARRIVED_NOT_INTERPRETED = "arrived but could not be interpreted"


class _Unresolved(Exception):
    def __init__(self, kind: str, detail: str):
        super().__init__(detail)
        self.kind = kind
        self.detail = detail


class WorkbookError(Exception):
    """A workbook that could not be turned into a table; the words name the file."""


@dataclass(frozen=True)
class DecodedTable:
    """What one payload decoded to, or why it did not."""

    filename: str
    status: str
    text: str | None
    format: str | None
    compression: str | None
    encoding: str | None
    source_checksum: str
    source_bytes: int
    decoded_bytes: int | None
    reason_kind: str | None = None
    reason: str | None = None
    decoder_version: int = DECODER_VERSION

    @property
    def ok(self) -> bool:
        return self.status == DECODED

    def provenance(self) -> dict:
        """The facts recorded with the decoded artifact."""
        return {
            "status": self.status,
            "source_checksum": self.source_checksum,
            "source_bytes": self.source_bytes,
            "decoded_bytes": self.decoded_bytes,
            "compression": self.compression,
            "encoding": self.encoding,
            "format": self.format,
            "decoder_version": self.decoder_version,
            "reason_kind": self.reason_kind,
            "reason": self.reason,
        }


def checksum(raw: bytes) -> str:
    return hashlib.sha256(raw or b"").hexdigest()


def decode_table(raw: bytes, filename: str = "", *, max_decompressed_bytes: int | None = None) -> DecodedTable:
    """Decode one payload into table text, or say exactly why not. Never raises."""
    name = filename or "the file"
    limit = max_decompressed_bytes if max_decompressed_bytes is not None else DEFAULT_MAX_DECOMPRESSED_BYTES
    raw = bytes(raw or b"")
    compression = None
    encoding = None
    fmt = None
    try:
        if not raw:
            raise _Unresolved(EMPTY, "it is empty")
        compression, data = _decompress(raw, limit)
        if not data:
            raise _Unresolved(EMPTY, "it is empty")
        if data.startswith(_OLE2):
            fmt = "xls"
            text = _workbook_text(data, name, reader=xls_to_tsv)
        elif data.startswith(_ZIP) and _is_workbook(data):
            fmt = "xlsx"
            text = _workbook_text(data, name, reader=xlsx_to_tsv)
        else:
            fmt = "text"
            encoding, text = _decode_text(data)
        if "\x00" in text:
            raise _Unresolved(NUL, "its decoded text holds NUL characters, which no table's text contains")
    except _Unresolved as exc:
        return DecodedTable(
            filename=name,
            status=UNRESOLVED,
            text=None,
            format=fmt,
            compression=compression,
            encoding=encoding,
            source_checksum=checksum(raw),
            source_bytes=len(raw),
            decoded_bytes=None,
            reason_kind=exc.kind,
            reason=f"{name} {ARRIVED_NOT_INTERPRETED}: {safe_excerpt(exc.detail, limit=400)}",
        )
    return DecodedTable(
        filename=name,
        status=DECODED,
        text=text,
        format=fmt,
        compression=compression,
        encoding=encoding,
        source_checksum=checksum(raw),
        source_bytes=len(raw),
        decoded_bytes=len(text.encode("utf-8")),
    )


def _decompress(raw: bytes, limit: int) -> tuple[str | None, bytes]:
    """``(compression, payload)``: the enclosed bytes, or the bytes unchanged when not compressed."""
    if raw.startswith(_GZIP):
        return "gzip", _gunzip(raw, limit)
    if raw.startswith(_BZIP2):
        return "bzip2", _stream(bz2.BZ2Decompressor(), raw, limit, "bzip2")
    if raw.startswith(_XZ):
        return "xz", _stream(lzma.LZMADecompressor(), raw, limit, "xz")
    if raw.startswith(_ZIP) and not _is_workbook(raw):
        return "zip", _unzip(raw, limit)
    return None, raw


def _gunzip(raw: bytes, limit: int) -> bytes:
    """Every gzip member, within the limit. Trailing zero padding is tolerated; anything else is corrupt."""
    out = bytearray()
    data = raw
    while data:
        if not data.startswith(_GZIP):
            if data.strip(b"\x00"):
                raise _Unresolved(CORRUPT, "its gzip stream is followed by bytes that are not gzip")
            break
        decompressor = zlib.decompressobj(16 + zlib.MAX_WBITS)
        try:
            out += decompressor.decompress(data, limit + 1 - len(out))
        except zlib.error as exc:
            raise _Unresolved(CORRUPT, f"its gzip stream is corrupt ({exc})") from exc
        if len(out) > limit or decompressor.unconsumed_tail:
            raise _Unresolved(TOO_LARGE, f"it decompresses to more than {limit} bytes, bioAF's limit")
        if not decompressor.eof:
            raise _Unresolved(CORRUPT, "its gzip stream is truncated")
        data = decompressor.unused_data
    return bytes(out)


def _stream(decompressor, raw: bytes, limit: int, label: str) -> bytes:
    try:
        out = decompressor.decompress(raw, max_length=limit + 1)
    except (OSError, EOFError, lzma.LZMAError, ValueError) as exc:
        raise _Unresolved(CORRUPT, f"its {label} stream is corrupt ({exc})") from exc
    if len(out) > limit or not decompressor.needs_input and not decompressor.eof:
        raise _Unresolved(TOO_LARGE, f"it decompresses to more than {limit} bytes, bioAF's limit")
    if not decompressor.eof:
        raise _Unresolved(CORRUPT, f"its {label} stream is truncated")
    return out


def _unzip(raw: bytes, limit: int) -> bytes:
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            members = [info for info in archive.infolist() if not info.is_dir()]
            if len(members) != 1:
                raise _Unresolved(
                    ARCHIVE, f"it is an archive holding {len(members)} files, and bioAF reads one table at a time"
                )
            if members[0].file_size > limit:
                raise _Unresolved(TOO_LARGE, f"it decompresses to more than {limit} bytes, bioAF's limit")
            with archive.open(members[0]) as handle:
                out = handle.read(limit + 1)
    except (zipfile.BadZipFile, OSError, EOFError, zlib.error) as exc:
        raise _Unresolved(CORRUPT, f"its zip archive is corrupt ({exc})") from exc
    if len(out) > limit:
        raise _Unresolved(TOO_LARGE, f"it decompresses to more than {limit} bytes, bioAF's limit")
    return out


def _is_workbook(data: bytes) -> bool:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            return any(name.startswith("xl/") for name in archive.namelist())
    except (zipfile.BadZipFile, OSError, EOFError):
        return False


def _decode_text(data: bytes) -> tuple[str, str]:
    """``(encoding, text)``, strictly. A byte-order mark names the encoding; otherwise only UTF-8 is read."""
    for mark, encoding in _MARKS:
        if data.startswith(mark):
            codec = encoding
            body = data if encoding == "utf-8-sig" else data[len(mark) :]
            try:
                return encoding, body.decode(codec)
            except UnicodeDecodeError as exc:
                raise _Unresolved(
                    ENCODING, f"its bytes are not valid {encoding.upper()} after its byte-order mark ({exc.reason})"
                ) from exc
    try:
        return "utf-8", data.decode("utf-8")
    except UnicodeDecodeError as exc:
        head = data[:4096]
        if _looks_like_wide_text(head):
            raise _Unresolved(
                ENCODING, "it looks like UTF-16 or UTF-32 text with no byte-order mark, so its encoding is uncertain"
            ) from exc
        if b"\x00" in head or _control_share(head) > 0.1:
            raise _Unresolved(
                BINARY, "it is not a text table or a workbook bioAF reads (its bytes look binary)"
            ) from exc
        raise _Unresolved(
            ENCODING, "its bytes are not valid UTF-8, and it carries no byte-order mark naming another encoding"
        ) from exc


def _looks_like_wide_text(head: bytes) -> bool:
    """NUL bytes in every other position of mostly printable text: wide text with its mark missing."""
    if len(head) < 8:
        return False
    evens, odds = head[0::2], head[1::2]
    for zeros, chars in ((odds, evens), (evens, odds)):
        if zeros.count(0) >= 0.9 * len(zeros) and _control_share(chars.replace(b"\x00", b"")) < 0.1:
            return True
    return False


def _control_share(data: bytes) -> float:
    if not data:
        return 0.0
    controls = sum(1 for b in data if b < 0x20 and b not in (0x09, 0x0A, 0x0D))
    return controls / len(data)


def _workbook_text(data: bytes, name: str, *, reader) -> str:
    try:
        return reader(data, name)
    except WorkbookError as exc:
        detail = str(exc)
        # The reason already begins with the file's name.
        if detail.startswith(f"{name} is "):
            detail = "it is " + detail[len(name) + 4 :]
        raise _Unresolved(UNREADABLE_WORKBOOK, detail) from exc


def xls_to_tsv(data: bytes, filename: str) -> str:
    """First worksheet of a legacy BIFF .xls as TSV.

    xlrd 2.x reads ONLY .xls: it deliberately dropped .xlsx support, which makes it the exact
    complement of openpyxl rather than an alternative to it. GSE213770's differential-methylation
    table is one of these.
    """
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=data)
    except Exception as exc:  # noqa: BLE001 - a corrupt workbook is a deposit problem, not a crash
        raise WorkbookError(f"{filename} is a legacy Excel file that could not be opened: {exc}") from exc

    if not book.nsheets:
        raise WorkbookError(f"{filename} is a legacy Excel file with no worksheets")
    sheet = book.sheet_by_index(0)

    lines: list[str] = []
    for r in range(sheet.nrows):
        cells = []
        for value in sheet.row_values(r):
            # xlrd returns every number as a float, so an integer count would render as "5.0" and the
            # matrix would be measured as normalized rather than as counts.
            if isinstance(value, float) and value.is_integer():
                cells.append(str(int(value)))
            else:
                cells.append("" if value is None else str(value))
        if any(c.strip() for c in cells):
            lines.append("\t".join(cells))
    if not lines:
        raise WorkbookError(f"{filename} is a legacy Excel file with no rows")
    return "\n".join(lines) + "\n"


def xlsx_to_tsv(data: bytes, filename: str) -> str:
    """First worksheet of an .xlsx as TSV."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - a corrupt workbook is a deposit problem, not a crash
        raise WorkbookError(f"{filename} is an Excel file that could not be opened: {exc}") from exc

    ws = wb.worksheets[0] if wb.worksheets else None
    if ws is None:
        raise WorkbookError(f"{filename} is an Excel file with no worksheets")

    lines: list[str] = []
    for row in ws.iter_rows(values_only=True):
        if row is None:
            continue
        cells = ["" if c is None else str(c) for c in row]
        if any(c.strip() for c in cells):
            lines.append("\t".join(cells))
    if not lines:
        raise WorkbookError(f"{filename} is an Excel file with no rows")
    return "\n".join(lines) + "\n"


def decode_prefix(raw: bytes, *, max_bytes: int) -> tuple[str | None, bool, str | None]:
    """``(text, truncated, error)`` for the first bytes of a file, for a bounded preview.

    The same order as ``decode_table`` over a prefix: a gzip stream is decompressed as far as it goes,
    a byte-order mark names the encoding, and an incomplete final character is dropped rather than
    replaced. Nothing here reaches a comparison; a preview only shows a header and a few rows.
    """
    import codecs

    data = bytes(raw or b"")
    truncated = True
    if data.startswith(_GZIP):
        decoder = zlib.decompressobj(wbits=zlib.MAX_WBITS | 32)
        try:
            data = decoder.decompress(data, max_bytes)
        except zlib.error:
            data = b""
        truncated = not decoder.eof
    encoding = "utf-8"
    for mark, name in _MARKS:
        if data.startswith(mark):
            encoding = name
            data = data if name == "utf-8-sig" else data[len(mark) :]
            break
    try:
        text = codecs.getincrementaldecoder(encoding)(errors="strict").decode(data, final=not truncated)
    except UnicodeDecodeError as exc:
        return None, truncated, f"the file's first bytes are not valid {encoding.upper()} ({exc.reason})"
    if "\x00" in text:
        return None, truncated, "the file's first bytes hold NUL characters, so its encoding is uncertain"
    return text, truncated, None


def safe_excerpt(value, *, limit: int = 200) -> str:
    """Diagnostic text with every control character escaped, so it can be logged and stored."""
    if isinstance(value, (bytes, bytearray)):
        value = bytes(value).decode("latin-1")
    text = str(value if value is not None else "")
    out = []
    for char in text:
        if char == "\t":
            out.append("\\t")
        elif char == "\n":
            out.append("\\n")
        elif char == "\r":
            out.append("\\r")
        elif ord(char) < 0x20 or ord(char) == 0x7F:
            out.append(f"\\x{ord(char):02x}")
        else:
            out.append(char)
    escaped = "".join(out)
    return escaped if len(escaped) <= limit else escaped[: max(0, limit - 3)] + "..."


def json_safe(value):
    """A payload PostgreSQL's JSONB accepts: NUL characters escaped in every string, nothing else changed.

    For diagnostics and error records only. Scientific content never needs this, because a table whose
    decoded text holds a NUL is never interpreted."""
    if isinstance(value, str):
        return value.replace("\x00", "\\x00")
    if isinstance(value, dict):
        return {json_safe(k) if isinstance(k, str) else k: json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [json_safe(v) for v in value]
    return value
